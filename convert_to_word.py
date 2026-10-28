import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def create_word_document():
    md_path = r"C:\Users\HP\.gemini\antigravity-ide\scratch\medismart-ai\PROJECT_REPORT.md"
    docx_path = r"C:\Users\HP\.gemini\antigravity-ide\scratch\medismart-ai\PROJECT_REPORT.docx"

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()

    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base Normal Style: Times New Roman 12pt
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    i = 0
    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal table_rows, doc
        if not table_rows:
            return
        
        # Parse table rows (filter out separator row like |---|---|)
        data_rows = []
        for r in table_rows:
            raw_cols = [c.strip() for c in r.split('|')]
            if len(raw_cols) >= 3:
                cols = raw_cols[1:-1]
                # Check if separator row
                if all(re.match(r'^:?-+:?$', c) for c in cols if c):
                    continue
                data_rows.append(cols)

        if not data_rows:
            table_rows = []
            return

        num_cols = max(len(r) for r in data_rows)
        tbl = doc.add_table(rows=len(data_rows), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = True

        for row_idx, row_data in enumerate(data_rows):
            for col_idx in range(num_cols):
                val = row_data[col_idx] if col_idx < len(row_data) else ""
                # Clean markdown formatting like **bold** or `code`
                clean_val = re.sub(r'\*\*(.*?)\*\*', r'\1', val)
                clean_val = re.sub(r'`(.*?)`', r'\1', clean_val)
                
                cell = tbl.cell(row_idx, col_idx)
                cell.text = clean_val
                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                run = p.runs[0] if p.runs else p.add_run(clean_val)
                run.font.name = 'Times New Roman'
                
                if row_idx == 0:
                    set_cell_background(cell, "1E3A8A")  # Deep blue header
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    run.font.size = Pt(10)
                else:
                    if row_idx % 2 == 1:
                        set_cell_background(cell, "F8FAFC")
                    else:
                        set_cell_background(cell, "FFFFFF")
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

        # Space after table
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(0)
        p_after.paragraph_format.space_after = Pt(4)
        table_rows = []

    def flush_code():
        nonlocal code_lines, doc
        if not code_lines:
            return
        
        code_text = "".join(code_lines)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.right_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        code_lines = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith("```"):
            if in_code_block:
                in_code_block = False
                flush_code()
            else:
                if in_table:
                    in_table = False
                    flush_table()
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Handle tables
        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_rows.append(stripped)
            i += 1
            continue
        else:
            if in_table:
                in_table = False
                flush_table()

        # Blank line
        if not stripped:
            i += 1
            continue

        # Horizontal rule or break
        if stripped in ["---", "___", "***", "<br>", "<br><br>"]:
            i += 1
            continue

        # Headings
        if stripped.startswith("# "):
            heading_text = stripped[2:].strip()
            # If start of Chapter, add page break
            if "CHAPTER " in heading_text.upper() or "BONAFIDE CERTIFICATE" in heading_text.upper() or "ACKNOWLEDGMENT" in heading_text.upper() or "ABSTRACT" in heading_text.upper() or "TABLE OF CONTENTS" in heading_text.upper():
                doc.add_page_break()
            
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ("II SEMESTER" in heading_text or "PHARMAAI" in heading_text or "MEDISMART" in heading_text or "CERTIFICATE" in heading_text or "ACKNOWLEDGMENT" in heading_text or "ABSTRACT" in heading_text or "TABLE OF CONTENTS" in heading_text or "CHAPTER " in heading_text) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(heading_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            i += 1
            continue

        if stripped.startswith("## "):
            heading_text = stripped[3:].strip()
            if any(term in heading_text.upper() for term in ["BONAFIDE CERTIFICATE", "ACKNOWLEDGMENT", "ABSTRACT"]):
                doc.add_page_break()
                
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if any(term in heading_text.upper() for term in ["BONAFIDE CERTIFICATE", "ACKNOWLEDGMENT", "ABSTRACT", "TABLE OF CONTENTS"]) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(heading_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            i += 1
            continue

        if stripped.startswith("### "):
            heading_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(heading_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            i += 1
            continue

        if stripped.startswith("#### "):
            heading_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(heading_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
            i += 1
            continue

        # Bullet list
        if stripped.startswith("* ") or stripped.startswith("- "):
            bullet_text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            # Simple bold parsing
            parts = re.split(r'(\*\*.*?\*\*)', bullet_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.name = 'Times New Roman'
                    r.font.bold = True
                else:
                    clean = re.sub(r'`(.*?)`', r'\1', part)
                    r = p.add_run(clean)
                    r.font.name = 'Times New Roman'
            i += 1
            continue

        # Numbered list
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            content = num_match.group(2)
            parts = re.split(r'(\*\*.*?\*\*)', content)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.name = 'Times New Roman'
                    r.font.bold = True
                else:
                    clean = re.sub(r'`(.*?)`', r'\1', part)
                    r = p.add_run(clean)
                    r.font.name = 'Times New Roman'
            i += 1
            continue

        # Normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15

        # Center alignment for title page text
        if any(term in stripped for term in ["Submitted in partial fulfilment", "MASTER OF COMPUTER APPLICATIONS", "Kristu Jayanti", "Bengaluru – 560077"]):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        parts = re.split(r'(\*\*.*?\*\*)', stripped)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                r.font.name = 'Times New Roman'
                r.font.bold = True
            else:
                clean = re.sub(r'`(.*?)`', r'\1', part)
                r = p.add_run(clean)
                r.font.name = 'Times New Roman'

        i += 1

    if in_table:
        flush_table()
    if in_code_block:
        flush_code()

    doc.save(docx_path)
    print(f"Successfully created: {docx_path}")

if __name__ == "__main__":
    create_word_document()
